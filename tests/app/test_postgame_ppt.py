"""Tests de `app/reports/postgame_ppt.py`: puntos destacados (reglas + LLM) y
generación del `.pptx` del informe "PPT para Paolo".

Sin red y sin Streamlit: la selección de estadísticas y la maquetación de
diapositivas son funciones puras sobre diccionarios, igual que las
herramientas del asistente (ver `tests/app/assistant/test_tools.py`) — solo
la fachada `generate_postgame_ppt` toca la base de datos, y esa la cubre
`tests/app/test_queries.py::test_game_player_report_*`.
"""
import io

from pptx import Presentation

from app.assistant.llm.base import LLMError, LLMResponse
from app.reports import postgame_ppt


def _row(**overrides):
    base = {
        "player_id": "howard", "player_name": "Markus Howard", "number": 0, "position": "Base",
        "photo_url": None, "photo_local_path": None,
        "minutes": 28.0, "pts": 12, "reb": 3, "ast": 4, "efg_pct": 50.0,
        "stl": 1, "tov": 2, "blk": 0, "blk_against": 0, "pf": 2, "pf_drawn": 2,
        "oreb": 1, "dreb": 2, "plus_minus": 3, "pir": 12, "dunks": 0,
        "ftm": 2, "fta": 2, "tpm": 1.0, "tpa": 3.0,
    }
    base.update(overrides)
    return base


class _FakeClient:
    """Doble de `LLMClient` (§ `llm/base.py`): devuelve `text` tal cual, sin red."""

    def __init__(self, text):
        self._text = text
        self.calls = []

    def chat(self, messages, tools, *, system, on_text=None):
        self.calls.append({"messages": messages, "tools": tools, "system": system})
        return LLMResponse(text=self._text)


class _BrokenClient:
    def chat(self, *args, **kwargs):
        raise LLMError("sin saldo")


# --------------------------------------------------------- reglas (fallback) --


def test_rule_based_highlights_flags_hot_three_point_shooting():
    highlights = postgame_ppt._rule_based_highlights(_row(tpm=5, tpa=8))
    assert any("triples" in text for text in highlights)


def test_rule_based_highlights_flags_cold_three_point_shooting():
    highlights = postgame_ppt._rule_based_highlights(_row(tpm=0, tpa=5))
    assert "0/5 en triples" in highlights


def test_rule_based_highlights_flags_many_turnovers():
    highlights = postgame_ppt._rule_based_highlights(_row(tov=6))
    assert any("pérdidas" in text for text in highlights)


def test_rule_based_highlights_flags_efficient_low_minutes():
    highlights = postgame_ppt._rule_based_highlights(_row(minutes=9.0, pts=10))
    assert any("solo" in text.lower() and "min" in text for text in highlights)


def test_rule_based_highlights_never_empty_without_notable_stats():
    """Un jugador sin nada que cruce ningún umbral se queda igualmente con una
    frase (minutos + puntos) — nunca una diapositiva sin texto."""
    row = _row(
        pts=6, reb=1, ast=0, tov=1, stl=0, blk=0, pf=1, pir=6, plus_minus=0,
        ftm=0, fta=0, tpm=0, tpa=0, minutes=18.0, efg_pct=40.0,
    )
    highlights = postgame_ppt._rule_based_highlights(row)
    assert len(highlights) == 1
    assert "puntos" in highlights[0]


def test_rule_based_highlights_caps_at_max_items():
    row = _row(tpm=6, tpa=9, reb=12, ast=8, tov=6, stl=4, blk=3, pf=6, pir=25, plus_minus=18)
    highlights = postgame_ppt._rule_based_highlights(row, max_items=4)
    assert len(highlights) == 4


def test_rule_based_highlights_tolerates_missing_box_extras():
    """Partido sin boxscore ampliado (`stl`/`tov`/`blk`/`pf`/`pir`/`plus_minus`
    en NULL, Fase 1 aún no cargada): no debe reventar, solo tiene menos de
    donde sacar algo destacado."""
    row = _row(stl=None, tov=None, blk=None, pf=None, pir=None, plus_minus=None, dunks=None)
    highlights = postgame_ppt._rule_based_highlights(row)
    assert highlights  # sigue habiendo al menos la línea básica o el triple


# --------------------------------------------------------------------- LLM --


def test_parse_llm_json_accepts_plain_json():
    result = postgame_ppt._parse_llm_json('{"howard": ["frase uno", "frase dos"]}')
    assert result == {"howard": ["frase uno", "frase dos"]}


def test_parse_llm_json_strips_code_fence():
    result = postgame_ppt._parse_llm_json('```json\n{"howard": ["frase uno"]}\n```')
    assert result == {"howard": ["frase uno"]}


def test_parse_llm_json_returns_none_on_garbage():
    assert postgame_ppt._parse_llm_json("no soy json, lo siento") is None


def test_parse_llm_json_drops_entries_with_empty_lists():
    result = postgame_ppt._parse_llm_json('{"howard": ["ok"], "moneke": []}')
    assert result == {"howard": ["ok"]}


def test_select_highlights_prefers_llm_result_per_player():
    rows = [_row(player_id="howard"), _row(player_id="moneke", player_name="Chima Moneke", number=95)]
    client = _FakeClient('{"howard": ["Partidazo desde el triple"]}')

    highlights = postgame_ppt.select_highlights(client, rows, {"rival": "Real Madrid"})

    assert highlights["howard"] == ["Partidazo desde el triple"]
    # 'moneke' no viene en la respuesta del LLM: se queda con el fallback por reglas.
    assert highlights["moneke"] == postgame_ppt._rule_based_highlights(rows[1])
    assert len(client.calls) == 1  # una sola llamada para todos los jugadores, no una por jugador


def test_select_highlights_falls_back_to_rules_when_llm_errors():
    rows = [_row()]
    highlights = postgame_ppt.select_highlights(_BrokenClient(), rows, {})
    assert highlights["howard"] == postgame_ppt._rule_based_highlights(rows[0])


def test_select_highlights_falls_back_to_rules_on_unparseable_response():
    rows = [_row()]
    highlights = postgame_ppt.select_highlights(_FakeClient("respuesta que no es JSON"), rows, {})
    assert highlights["howard"] == postgame_ppt._rule_based_highlights(rows[0])


def test_select_highlights_without_client_uses_rules_only():
    rows = [_row()]
    highlights = postgame_ppt.select_highlights(None, rows, {})
    assert highlights["howard"] == postgame_ppt._rule_based_highlights(rows[0])


# ----------------------------------------------------------------- .pptx --


def test_build_postgame_ppt_has_one_title_slide_plus_one_per_player():
    rows = [_row(player_id="howard"), _row(player_id="moneke", player_name="Chima Moneke", number=95)]
    highlights = {"howard": ["18 puntos"], "moneke": ["10 rebotes"]}

    data = postgame_ppt.build_postgame_ppt(rows, highlights, {"subtitle": "vs Real Madrid"})
    prs = Presentation(io.BytesIO(data))

    assert len(prs.slides) == 1 + len(rows)


def test_build_postgame_ppt_falls_back_to_initials_badge_without_photo():
    """Sin foto (ni local ni remota) la diapositiva no revienta: cae al badge
    de iniciales, igual que `avatar.avatar_html` en el resto de la interfaz."""
    rows = [_row(photo_local_path=None, photo_url=None)]

    data = postgame_ppt.build_postgame_ppt(rows, {"howard": ["18 puntos"]}, {"subtitle": ""})

    assert len(data) > 0
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2


def test_build_postgame_ppt_player_slide_contains_the_chosen_highlights():
    rows = [_row(player_id="howard", player_name="Markus Howard")]
    highlights = {"howard": ["9/11 en triples", "5 pérdidas"]}

    data = postgame_ppt.build_postgame_ppt(rows, highlights, {"subtitle": ""})
    prs = Presentation(io.BytesIO(data))
    player_slide = list(prs.slides)[1]
    all_text = "\n".join(
        shape.text_frame.text for shape in player_slide.shapes if shape.has_text_frame
    )

    assert "Markus Howard" in all_text
    assert "9/11 en triples" in all_text
    assert "5 pérdidas" in all_text
