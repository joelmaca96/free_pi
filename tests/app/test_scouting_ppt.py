"""Tests de `app/reports/scouting_ppt.py`: dossier de scouting del rival.

Mismo criterio que `test_postgame_ppt.py` (léase su docstring): sin red y sin
Streamlit, sobre funciones puras y datos de mentira — `generate_scouting_ppt`
es la única pieza que toca la base de datos y no se prueba aquí (sus
consultas ya están una a una en `test_queries.py`/`test_queries_assistant.py`
cuando existen, y volver a montar un fixture de BD completo con percentiles,
tiros y quintetos solo para ejercitar el ensamblado sería probar dos veces lo
mismo con más aparato). `build_scouting_ppt` sí es pura y es la que se prueba
de punta a punta.
"""
import io

import pandas as pd
import pytest
from pptx import Presentation

from app.analytics import shot_quality
from app.assistant.llm.base import LLMError, LLMResponse
from app.reports import scouting_ppt


def _player_row(**overrides):
    base = {
        "id": "howard", "name": "Marcus Howard", "number": 0, "position": "Base",
        "photo_url": None, "photo_local_path": None,
        "gp": 20, "min_avg": 28.0, "pts_avg": 15.0, "reb_avg": 3.0, "ast_avg": 4.5, "efg_pct": 52.0,
    }
    base.update(overrides)
    return base


def _ctx(**overrides):
    base = {
        "rival_name": "Real Madrid",
        "is_home": True,
        "competition": "Euroliga",
        "fecha": "2026-10-01",
        "h2h_summary": None,
        "rival_record": [],
        "is_fallback_season": False,
        "scouting_season_label": "2025-2026",
        "style_row": None,
        "quarters_df": pd.DataFrame(columns=["quarter", "avg_points_for", "avg_points_against", "gp"]),
        "top_player": None,
        "attack_diff": None,
        "defense_diff": None,
        "win_threshold_cards": [],
    }
    base.update(overrides)
    return base


class _FakeClient:
    """Doble de `LLMClient`: devuelve `text` tal cual, sin red."""

    def __init__(self, text):
        self._text = text
        self.calls = []

    def chat(self, messages, tools, *, system, on_text=None):
        self.calls.append({"messages": messages, "tools": tools, "system": system})
        return LLMResponse(text=self._text)


class _BrokenClient:
    def chat(self, *args, **kwargs):
        raise LLMError("sin saldo")


def _all_text(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


def _all_tables_text(slide) -> str:
    bits = []
    for shape in slide.shapes:
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    bits.append(cell.text)
    return "\n".join(bits)


# ------------------------------------------------------------- mapa de tiro --


def _zone_profile_row(zone_label, volume, made, pps, league_pps):
    diff_pps = pps - league_pps
    return {
        "zone_label": zone_label, "volume": volume, "made": made,
        "fg_pct": 100.0 * made / volume, "pps": pps, "league_pps": league_pps, "diff_pps": diff_pps,
    }


def test_coarsen_zone_profile_merges_the_wing_2pt_and_3pt_split():
    """"Ala izq. (2)"/"Ala izq. (3)" (partidas por el arco) se funden en una sola "Ala izq."."""
    zone_df = pd.DataFrame([
        _zone_profile_row("Ala izq. (2)", 100, 68, 1.36, 1.30),
        _zone_profile_row("Ala izq. (3)", 50, 20, 1.20, 1.10),
        _zone_profile_row("Pintura", 200, 120, 1.20, 1.15),
    ])
    coarse = scouting_ppt._coarsen_zone_profile(zone_df)

    assert sorted(coarse["zone_label"]) == ["Ala izq.", "Pintura"]
    wing = coarse.loc[coarse["zone_label"] == "Ala izq."].iloc[0]
    assert wing["volume"] == 150
    assert wing["made"] == 88
    # PPS conjunto: puntos totales / volumen total, no la media simple de los dos.
    expected_pps = (100 * 1.36 + 50 * 1.20) / 150
    assert wing["pps"] == pytest.approx(expected_pps)


def test_coarsen_zone_profile_of_empty_input_is_empty():
    empty = pd.DataFrame(columns=["zone_label", "volume", "made", "fg_pct", "pps", "league_pps", "diff_pps"])
    assert scouting_ppt._coarsen_zone_profile(empty).empty


def test_zone_fill_color_is_neutral_without_data():
    assert scouting_ppt._zone_fill_color(None) == scouting_ppt._ZONE_EMPTY
    assert scouting_ppt._zone_fill_color(float("nan")) == scouting_ppt._ZONE_EMPTY


def test_zone_fill_color_saturates_at_the_cap_in_both_directions():
    better = scouting_ppt._zone_fill_color(scouting_ppt._ZONE_DIFF_CAP * 5)  # muy por encima del cap
    worse = scouting_ppt._zone_fill_color(-scouting_ppt._ZONE_DIFF_CAP * 5)
    assert better == scouting_ppt._lerp_rgb(scouting_ppt._ZONE_NEUTRAL, scouting_ppt._ZONE_HIGH, 1.0)
    assert worse == scouting_ppt._lerp_rgb(scouting_ppt._ZONE_NEUTRAL, scouting_ppt._ZONE_LOW, 1.0)


def _court_zones_df():
    """Subconjunto realista de `queries.court_zones`: solo rectángulos reales (ver `court_zones` en
    la base de datos) — incluye las filas degeneradas (`x_min == x_max`) que el mapa debe ignorar."""
    return pd.DataFrame([
        {"id": 1, "label": "Pintura", "x_min": 195.0, "x_max": 305.0, "y_min": 300.0, "y_max": 455.0},
        {"id": 2, "label": "Ala izq.", "x_min": 0.0, "x_max": 195.0, "y_min": 170.0, "y_max": 380.0},
        {"id": 3, "label": "Ala der.", "x_min": 305.0, "x_max": 500.0, "y_min": 170.0, "y_max": 380.0},
        {"id": 10, "label": "Mate", "x_min": 250.0, "x_max": 250.0, "y_min": 455.0, "y_max": 455.0},
        {"id": 14, "label": "Ala izq. (2)", "x_min": 100.0, "x_max": 100.0, "y_min": 300.0, "y_max": 300.0},
    ])


def test_build_scouting_ppt_shot_quality_slide_draws_zone_map_when_geometry_present():
    attack = pd.DataFrame([_zone_profile_row("Pintura", 200, 130, 1.30, 1.10)])
    data = _build(zones_df=_court_zones_df(), attack_zone_profile=attack, defense_zone_profile=attack)
    prs = Presentation(io.BytesIO(data))
    slide = list(prs.slides)[2]  # portada, identidad, CALIDAD DE TIRO
    # 2 fondos de cancha + 3 zonas reales (Pintura/Ala izq./Ala der.) x 2 diagramas, más título,
    # subtítulo, viñetas, 2 cabeceras y el pie: bastante más que sin mapa (title+subtitle+body=3).
    assert len(slide.shapes) >= 14
    assert "Ataque (genera)" in _all_text(slide)
    assert "Defensa (concede)" in _all_text(slide)
    assert shot_quality.format_pps(1.30) in _all_text(slide)  # el PPS de Pintura, dibujado en su rectángulo


def test_build_scouting_ppt_shot_quality_slide_degrades_without_geometry():
    """Sin `zones_df` (o vacío), la diapositiva se queda solo con las viñetas — nunca revienta."""
    data = _build(shot_quality_bullets=["Ataque por encima de lo esperado."])
    prs = Presentation(io.BytesIO(data))
    slide = list(prs.slides)[2]
    assert "Ataque por encima de lo esperado." in _all_text(slide)
    assert "Ataque (genera)" not in _all_text(slide)


# --------------------------------------------- jugadores: reglas (fallback) --


def test_player_highlights_flags_high_scoring_percentile():
    row = _player_row(pts_avg=22.0, pts_pct=0.95)
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert any("anotadora" in text for text in highlights)
    assert any("95" in text for text in highlights)


def test_player_highlights_flags_low_scoring_percentile():
    row = _player_row(pts_avg=2.0, pts_pct=0.05)
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert any("Aporta poco" in text for text in highlights)


def test_player_highlights_flags_turnover_prone_player():
    """`tov_pct` viene invertido en la vista: percentil BAJO = pierde mucho balón."""
    row = _player_row(tov_avg=4.2, tov_pct=0.05)
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert any("Cuida poco el balón" in text for text in highlights)


def test_player_highlights_ignores_efficiency_extremes_on_low_usage():
    """Un jugador que solo tira una vez y mete no debe salir como "muy eficiente"."""
    row = _player_row(pts_avg=2.0, efg_pct=100.0, efg_pct_pct=0.99)
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert not any("eficiente" in text for text in highlights)


def test_player_highlights_never_empty_without_extreme_percentiles():
    row = _player_row(pts_pct=0.5, reb_pct=0.5, ast_pct=0.5)
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert len(highlights) == 1
    assert "pts/partido" in highlights[0]


def test_player_highlights_caps_at_max_items():
    row = _player_row(
        pts_pct=0.95, reb_pct=0.95, ast_pct=0.95, stl_pct=0.95, blk_pct=0.95,
        pir_pct=0.95, stl_avg=3.0, blk_avg=2.0, pir_avg=25.0,
    )
    highlights = scouting_ppt._rule_based_player_highlights(row, max_items=3)
    assert len(highlights) == 3


def test_player_highlights_tolerates_missing_percentiles():
    """Jugador sin fila en `player_percentiles` (menos de 5 partidos): no revienta."""
    row = _player_row()
    highlights = scouting_ppt._rule_based_player_highlights(row)
    assert highlights  # cae a la línea básica


# ------------------------------------------------------------- jugadores: LLM --


def test_parse_llm_json_dict_accepts_plain_json():
    result = scouting_ppt._parse_llm_json_dict('{"howard": ["frase uno"]}', max_items=4)
    assert result == {"howard": ["frase uno"]}


def test_parse_llm_json_dict_returns_none_on_garbage():
    assert scouting_ppt._parse_llm_json_dict("no soy json", max_items=4) is None


def test_select_player_highlights_prefers_llm_result_per_player():
    rows = [_player_row(id="howard"), _player_row(id="moneke", name="Chima Moneke")]
    client = _FakeClient('{"howard": ["Tira mucho desde la esquina derecha"]}')

    highlights = scouting_ppt.select_player_highlights(client, rows, "Real Madrid")

    assert highlights["howard"] == ["Tira mucho desde la esquina derecha"]
    assert highlights["moneke"] == scouting_ppt._rule_based_player_highlights(rows[1])
    assert len(client.calls) == 1  # una sola llamada para todos los jugadores clave


def test_select_player_highlights_falls_back_to_rules_when_llm_errors():
    rows = [_player_row()]
    highlights = scouting_ppt.select_player_highlights(_BrokenClient(), rows, "Real Madrid")
    assert highlights["howard"] == scouting_ppt._rule_based_player_highlights(rows[0])


def test_select_player_highlights_without_client_uses_rules_only():
    rows = [_player_row()]
    highlights = scouting_ppt.select_player_highlights(None, rows, "Real Madrid")
    assert highlights["howard"] == scouting_ppt._rule_based_player_highlights(rows[0])


# ------------------------------------------------ claves del partido: reglas --


def test_game_keys_flags_high_pace_percentile():
    ctx = _ctx(style_row=pd.Series({"pace": 78.3, "pace_pct": 0.92, "ortg": 110.0, "drtg": 105.0, "net_rating": 5.0}))
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("juega rápido" in text for text in keys)


def test_game_keys_flags_weak_defense_as_exploitable():
    ctx = _ctx(style_row=pd.Series({"pace": 70.0, "ortg": 108.0, "drtg": 112.0, "drtg_pct": 0.1, "net_rating": -3.0}))
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("atacable" in text for text in keys)


def test_game_keys_includes_head_to_head_when_enough_history():
    ctx = _ctx(h2h_summary={"wins": 8, "losses": 2, "total": 10})
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("Cara a cara histórico" in text for text in keys)


def test_game_keys_ignores_head_to_head_with_too_little_history():
    ctx = _ctx(h2h_summary={"wins": 1, "losses": 0, "total": 1})
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert not any("Cara a cara" in text for text in keys)


def test_game_keys_flags_top_scorer_threat():
    ctx = _ctx(top_player={"name": "Facundo Campazzo", "pts_avg": 18.0, "pts_pct": 0.9})
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("Campazzo" in text for text in keys)


def test_game_keys_flags_shot_quality_diff():
    ctx = _ctx(attack_diff=0.12)
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("ataque" in text.lower() and "PPS" in text for text in keys)


def test_game_keys_includes_win_threshold_objective():
    """Propuesta 09 §6: los umbrales de victoria, ya ajustados al rival, alimentan las claves."""
    card = {
        "label": "rebote ofensivo", "higher_is_better": True,
        "display_threshold": 32.0, "adjusted_threshold": 34.0,
        "is_rival_adjusted": True, "separation": 18.0,
    }
    ctx = _ctx(rival_name="Real Madrid", win_threshold_cards=[card])
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any(
        "Objetivo del partido" in text and "34,0%" in text and "Real Madrid" in text for text in keys
    )


def test_game_keys_without_win_threshold_cards_still_works():
    """Sin ninguna fuente de claves (ctx "vacío"), la respuesta correcta es una lista vacía,
    no un error — `_build_keys_slide` cae al texto de "sin claves" con eso."""
    assert scouting_ppt._rule_based_game_keys(_ctx(win_threshold_cards=[])) == []


def test_game_keys_flags_uneven_quarter_split():
    quarters = pd.DataFrame({
        "quarter": [1, 2, 3, 4],
        "avg_points_for": [22.0, 20.0, 15.0, 18.0],
        "avg_points_against": [16.0, 19.0, 22.0, 17.0],
        "gp": [10, 10, 10, 10],
    })
    ctx = _ctx(quarters_df=quarters)
    keys = scouting_ppt._rule_based_game_keys(ctx)
    assert any("Reparto por cuarto desigual" in text for text in keys)


def test_game_keys_never_empty_is_allowed_but_bounded_at_max_items():
    ctx = _ctx(
        style_row=pd.Series({
            "pace": 78.0, "pace_pct": 0.95, "ortg": 115.0, "ortg_pct": 0.9,
            "drtg": 95.0, "drtg_pct": 0.95, "net_rating": 12.0, "net_rating_pct": 0.95,
        }),
        h2h_summary={"wins": 3, "losses": 7, "total": 10},
        top_player={"name": "X", "pts_avg": 20.0, "pts_pct": 0.9},
        attack_diff=0.2,
        defense_diff=-0.2,
    )
    keys = scouting_ppt._rule_based_game_keys(ctx, max_items=5)
    assert len(keys) == 5


def test_select_game_keys_prefers_llm_result():
    ctx = _ctx()
    client = _FakeClient('["clave uno", "clave dos", "clave tres", "clave cuatro", "clave cinco"]')
    keys = scouting_ppt.select_game_keys(client, ctx)
    assert keys == ["clave uno", "clave dos", "clave tres", "clave cuatro", "clave cinco"]


def test_select_game_keys_falls_back_to_rules_when_llm_errors():
    ctx = _ctx(h2h_summary={"wins": 5, "losses": 5, "total": 10})
    keys = scouting_ppt.select_game_keys(_BrokenClient(), ctx)
    assert keys == scouting_ppt._rule_based_game_keys(ctx)


def test_select_game_keys_falls_back_on_unparseable_response():
    ctx = _ctx(h2h_summary={"wins": 5, "losses": 5, "total": 10})
    keys = scouting_ppt.select_game_keys(_FakeClient("no soy json"), ctx)
    assert keys == scouting_ppt._rule_based_game_keys(ctx)


def test_select_game_keys_without_client_uses_rules_only():
    ctx = _ctx(h2h_summary={"wins": 5, "losses": 5, "total": 10})
    keys = scouting_ppt.select_game_keys(None, ctx)
    assert keys == scouting_ppt._rule_based_game_keys(ctx)


# ------------------------------------------------------------------- .pptx --


def _build(ctx=None, style_df=None, player_rows=None, player_highlights=None,
           lineups_df=None, shot_quality_bullets=None, game_keys=None,
           zones_df=None, attack_zone_profile=None, defense_zone_profile=None):
    return scouting_ppt.build_scouting_ppt(
        ctx if ctx is not None else _ctx(),
        style_df if style_df is not None else pd.DataFrame(),
        player_rows if player_rows is not None else [],
        player_highlights if player_highlights is not None else {},
        lineups_df if lineups_df is not None else pd.DataFrame(),
        shot_quality_bullets if shot_quality_bullets is not None else [],
        game_keys if game_keys is not None else [],
        zones_df=zones_df, attack_zone_profile=attack_zone_profile, defense_zone_profile=defense_zone_profile,
    )


def test_build_scouting_ppt_has_exactly_six_slides_with_one_player():
    data = _build(player_rows=[_player_row()], player_highlights={"howard": ["18 puntos de media"]})
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6


def test_build_scouting_ppt_adds_one_slide_per_extra_player():
    rows = [_player_row(id="howard"), _player_row(id="moneke", name="Chima Moneke")]
    data = _build(player_rows=rows, player_highlights={"howard": ["x"], "moneke": ["y"]})
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 7  # 6 fijas + 1 jugador extra


def test_build_scouting_ppt_degrades_gracefully_with_no_data_at_all():
    """Rival sin ningún dato cargado todavía: nunca debe reventar, y cada
    diapositiva cae a su texto de "sin datos" en vez de quedar en blanco."""
    data = _build(ctx=_ctx(rival_name="Recién Ascendido"))
    assert len(data) > 0
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6  # portada + identidad + tiro + jugadores(fallback) + quintetos + claves
    slides = list(prs.slides)
    assert "Recién Ascendido" in _all_text(slides[0])
    assert "no llega a los 5 partidos" in _all_text(slides[1])
    assert "Sin plantilla" in _all_text(slides[3])


def test_build_scouting_ppt_cover_slide_shows_fallback_season_warning():
    ctx = _ctx(is_fallback_season=True, scouting_season_label="2024-2025")
    data = _build(ctx=ctx)
    prs = Presentation(io.BytesIO(data))
    text = _all_text(list(prs.slides)[0])
    assert "2024-2025" in text
    assert "no ha jugado aún" in text


def test_build_scouting_ppt_identity_slide_has_percentile_table():
    style_df = pd.DataFrame([{
        "competition": "ACB", "competition_id": 1, "gp": 20, "league_teams": 18,
        "pace": 78.3, "pace_pct": 0.92, "ortg": 112.0, "ortg_pct": 0.7,
        "drtg": 108.0, "drtg_pct": 0.3, "net_rating": 4.0, "net_rating_pct": 0.6,
        "efg_pct": 54.0, "efg_pct_pct": 0.5, "ts_pct": 57.0, "ts_pct_pct": 0.5,
    }])
    data = _build(style_df=style_df)
    prs = Presentation(io.BytesIO(data))
    table_text = _all_tables_text(list(prs.slides)[1])
    assert "Ritmo" in table_text
    assert "78.3" in table_text
    assert "92" in table_text  # percentil de ritmo


def test_build_scouting_ppt_player_slide_contains_name_and_highlights():
    rows = [_player_row(name="Marcus Howard")]
    data = _build(player_rows=rows, player_highlights={"howard": ["Tira el 47% desde la esquina derecha"]})
    prs = Presentation(io.BytesIO(data))
    player_slide = list(prs.slides)[3]
    text = _all_text(player_slide)
    assert "Marcus Howard" in text
    assert "Tira el 47% desde la esquina derecha" in text


def test_build_scouting_ppt_lineups_slide_has_table_when_data_present():
    lineups_df = pd.DataFrame([
        {"jugadores": "A · B · C · D · E", "minutes": 120.5, "plus_minus": 15, "stints": 8, "total_combos": 40},
    ])
    data = _build(lineups_df=lineups_df)
    prs = Presentation(io.BytesIO(data))
    table_text = _all_tables_text(list(prs.slides)[4])
    assert "A · B · C · D · E" in table_text


def test_build_scouting_ppt_keys_slide_contains_the_chosen_keys():
    data = _build(game_keys=["Cuidado con su ritmo alto", "Explota el rebote ofensivo"])
    prs = Presentation(io.BytesIO(data))
    text = _all_text(list(prs.slides)[5])
    assert "Cuidado con su ritmo alto" in text
    assert "Explota el rebote ofensivo" in text
