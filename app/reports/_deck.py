"""Piezas de maquetación de `.pptx` compartidas entre los informes descargables.

Extraído de `postgame_ppt.py` (el primero en existir, "PPT para Paolo") cuando
`scouting_ppt.py` (dossier de scouting del rival, propuesta
`doc/features/propuestas/03_dossier_scouting_rival.md`) necesitó exactamente
las mismas piezas: la marca (`ACCENT`/`DARK`/`MUTED`/`WHITE`), el formato
16:9, y los tipos de diapositiva que se repiten en los dos informes — portada
con banda de color, foto+viñetas (una por jugador) y ahora también tabla
nativa y lista de viñetas a pantalla completa (quintetos, claves del
partido...).

Nada de este módulo sabe de baloncesto: todo lo que recibe es texto y
números ya calculados por quien llama (`postgame_ppt.py`/`scouting_ppt.py`),
nunca un DataFrame ni un id de la base de datos — eso mantiene la
maquetación reutilizable sin acoplarla a un esquema de consulta concreto.
"""
import io
import os
from typing import Optional, Sequence

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# `components.avatar` se importa igual venga el paquete por `app.*` (pytest,
# raíz en sys.path) o por su nombre corto (Streamlit, `app/` en sys.path) —
# mismo patrón que `assistant/tools/context.py`.
try:  # pragma: no cover - depende de cómo se arranque el proceso, no de la lógica
    from app.components import avatar
except ImportError:  # pragma: no cover
    from components import avatar

ACCENT = RGBColor(0x00, 0x83, 0x00)  # verde Baskonia — mismo accent que el resto de la interfaz
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_WIDTH_IN = 13.333  # 16:9
SLIDE_HEIGHT_IN = 7.5


def new_presentation() -> Presentation:
    """`Presentation` en blanco, ya en 16:9."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    return prs


def strip_shape_chrome(shape) -> None:
    """Sin borde ni sombra por defecto — sin esto cada rectángulo sale con una
    sombra gris que no pega con el resto, plano, de la interfaz."""
    shape.line.fill.background()
    shape.shadow.inherit = False


def blank_slide(prs: Presentation):
    """Diapositiva nueva con el layout en blanco (índice 6 en toda plantilla por defecto)."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_band_title_slide(
    prs: Presentation, *, title: str, subtitle: str = "", footer: str = "", crest_path: Optional[str] = None
):
    """Portada: banda de color arriba con escudo + título + subtítulo, nota al pie debajo."""
    slide = blank_slide(prs)

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(2.3))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT
    strip_shape_chrome(band)

    if crest_path and os.path.exists(crest_path):
        slide.shapes.add_picture(crest_path, Inches(0.6), Inches(0.5), height=Inches(1.3))

    title_box = slide.shapes.add_textbox(Inches(2.2), Inches(0.45), Inches(10.5), Inches(1.1))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(32), True, WHITE

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(2.2), Inches(1.4), Inches(10.5), Inches(0.7))
        p = subtitle_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size, p.font.color.rgb = Pt(18), WHITE

    if footer:
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(12.0), Inches(1.0))
        tf = footer_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = footer
        p.font.size, p.font.italic, p.font.color.rgb = Pt(14), True, MUTED
    return slide


def add_photo_bullets_slide(
    prs: Presentation,
    *,
    name: str,
    subtitle: str = "",
    stat_line: str = "",
    bullets: Sequence[str],
    photo_local_path=None,
    photo_url=None,
    fallback_text: str = "Sin nada destacado que señalar.",
):
    """Foto (o badge de iniciales) a la izquierda, cabecera + línea de cifras + viñetas a la derecha.

    Layout compartido por la diapositiva de jugador del informe de partido
    (`postgame_ppt.py`, una del PARTIDO) y la del dossier de scouting
    (`scouting_ppt.py`, una de la TEMPORADA) — el hueco visual es idéntico,
    solo cambia qué texto compone cada llamador.

    `photo_url` hoy no se usa para incrustar nada (python-pptx no descarga
    URLs remotas) — se acepta solo para que la firma sea simétrica con
    `components/avatar.py::avatar_html` y quede claro en el sitio de llamada
    que, a diferencia de la interfaz web, aquí una foto sin copia local
    (`photo_local_path`) cae directa al badge de iniciales.
    """
    slide = blank_slide(prs)

    photo_w, photo_h = Inches(3.6), Inches(5.4)  # ratio 2:3, igual que `avatar.player_avatar_html`
    left, top = Inches(0.5), Inches(1.1)

    raw = avatar.local_photo_bytes(photo_local_path) if pd.notna(photo_local_path) else None
    if raw is not None:
        slide.shapes.add_picture(io.BytesIO(raw), left, top, width=photo_w, height=photo_h)
    else:
        # Mismo badge de respaldo que la interfaz (`avatar.avatar_html`): sin
        # foto real, iniciales sobre verde Baskonia — nunca un hueco vacío.
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, photo_w, photo_h)
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT
        strip_shape_chrome(badge)
        tf = badge.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = avatar.initials(name)
        p.font.size, p.font.bold, p.font.color.rgb = Pt(96), True, WHITE
        p.alignment = PP_ALIGN.CENTER

    header = slide.shapes.add_textbox(Inches(4.5), Inches(0.5), Inches(8.3), Inches(1.1))
    tf = header.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = name
    p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, DARK
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size, p2.font.color.rgb = Pt(16), MUTED

    if stat_line:
        stat_box = slide.shapes.add_textbox(Inches(4.5), Inches(1.55), Inches(8.3), Inches(0.5))
        p = stat_box.text_frame.paragraphs[0]
        p.text = stat_line
        p.font.size, p.font.color.rgb = Pt(15), MUTED

    body = slide.shapes.add_textbox(Inches(4.5), Inches(2.35), Inches(8.3), Inches(4.3))
    tf = body.text_frame
    tf.word_wrap = True
    items = list(bullets) or [fallback_text]
    for index, text in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"●  {text}"
        p.font.size, p.font.color.rgb = Pt(22), DARK
        p.space_after = Pt(16)
    return slide


def add_bullets_slide(
    prs: Presentation,
    *,
    title: str,
    bullets: Sequence[str],
    subtitle: str = "",
    fallback_text: str,
    body_left: float = 0.6,
    body_width: float = 12.1,
):
    """Diapositiva de solo texto: título + lista de viñetas grandes (p.ej. "Claves del partido").

    Args:
        body_left/body_width: posición y ancho (en pulgadas) del subtítulo y
            del cuerpo de viñetas — a toda página por defecto (12.1", igual
            que antes de que existiera este parámetro). Un llamador que
            quiera dibujar algo MÁS en la diapositiva (p.ej. el mapa de tiro
            de `scouting_ppt._build_shot_quality_slide`) puede estrechar el
            cuerpo y añadir sus propias formas en el hueco libre sobre el
            `slide` devuelto, sin reimplementar el título/subtítulo/viñetas.
    """
    slide = blank_slide(prs)

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(30), True, ACCENT

    top_in = 1.2
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(body_left), Inches(top_in), Inches(body_width), Inches(0.5))
        p = subtitle_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size, p.font.color.rgb = Pt(16), MUTED
        top_in = 1.85

    body = slide.shapes.add_textbox(
        Inches(body_left), Inches(top_in), Inches(body_width), Inches(SLIDE_HEIGHT_IN - top_in - 0.4)
    )
    tf = body.text_frame
    tf.word_wrap = True
    items = list(bullets) or [fallback_text]
    for index, text in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"●  {text}"
        p.font.size, p.font.color.rgb = Pt(22), DARK
        p.space_after = Pt(18)
    return slide


def add_table_slide(
    prs: Presentation,
    *,
    title: str,
    columns: Sequence[str],
    rows: Sequence[Sequence],
    subtitle: str = "",
    footer: str = "",
    col_widths: Optional[Sequence[float]] = None,
):
    """Título + tabla NATIVA de pptx — no una imagen, así que sigue siendo editable
    y movible dentro de PowerPoint (§4 de la propuesta de dossier de scouting).

    Args:
        columns: cabeceras de columna, ya en el idioma final.
        rows: una lista por fila, cada celda ya convertida a texto (aquí no
            se sabe si una celda es un porcentaje, un entero o un nombre —
            esa decisión es de quien llama).
        col_widths: anchos en pulgadas, mismo orden que `columns`; si no se
            da, python-pptx los reparte iguales.
    """
    slide = blank_slide(prs)

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.7))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, ACCENT

    top_in = 1.05
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.6), Inches(top_in), Inches(12.1), Inches(0.45))
        p = subtitle_box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size, p.font.color.rgb = Pt(15), MUTED
        top_in = 1.55

    n_rows, n_cols = len(rows) + 1, len(columns)
    table_height_in = min(0.5 * n_rows, 5.6)
    shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.6), Inches(top_in), Inches(12.1), Inches(table_height_in)
    )
    table = shape.table
    if col_widths:
        for index, width in enumerate(col_widths):
            table.columns[index].width = Inches(width)

    for col_index, label in enumerate(columns):
        cell = table.cell(0, col_index)
        cell.text = str(label)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size, paragraph.font.bold, paragraph.font.color.rgb = Pt(13), True, WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT

    for row_index, row in enumerate(rows, start=1):
        # Cebra suave (par/impar) para que una tabla de 8-10 filas se siga
        # leyendo fila a fila sin regla debajo del dedo.
        row_fill = WHITE if row_index % 2 else RGBColor(0xEF, 0xEF, 0xEF)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size, paragraph.font.color.rgb = Pt(12), DARK
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_fill

    if footer:
        footer_box = slide.shapes.add_textbox(
            Inches(0.6), Inches(top_in + table_height_in + 0.15), Inches(12.1), Inches(0.6)
        )
        tf = footer_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = footer
        p.font.size, p.font.italic, p.font.color.rgb = Pt(12), True, MUTED
    return slide
